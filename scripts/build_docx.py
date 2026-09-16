#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成《大学物理 AI 智慧教学系统》演示文档 docx。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(BASE)
SHOTS = os.path.join(ROOT, "docs", "screenshots")
DIAGRAMS = os.path.join(ROOT, "docs", "diagrams")
OUT = os.path.join(ROOT, "docs", "物理AI教学系统-演示文档.docx")

doc = Document()

# 默认中文字体
style = doc.styles["Normal"]
style.font.name = "PingFang SC"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")

PRIMARY = RGBColor(0x03, 0x69, 0xA1)
DARK = RGBColor(0x0F, 0x1A, 0x2C)
MUTED = RGBColor(0x4A, 0x58, 0x69)


def set_east_asia(run, name="PingFang SC"):
    run.font.name = name
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_east_asia(run)
        run.font.color.rgb = DARK
    return h


def add_para(text, bold=False, size=11, color=None, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_east_asia(run)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if align:
        p.alignment = align
    return p


def add_image(path, width_cm=16.5, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        set_east_asia(run)
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


# ============ 封面 ============
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("大学物理 AI 智慧教学系统"); set_east_asia(r); r.font.size = Pt(30); r.bold = True; r.font.color.rgb = PRIMARY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("功能闭环演示文档"); set_east_asia(r); r.font.size = Pt(20); r.font.color.rgb = DARK
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("—— 诊断 → 项目匹配 → 任务闯关 → 虚拟实验 → 学习报告 ——"); set_east_asia(r); r.font.size = Pt(12); r.font.color.rgb = MUTED
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("适用对象：学生 / 教师（双角色）\n演示账号：student / teacher，统一密码 dunzhongwan\n技术栈：Next.js 16 + React 19 + Tailwind CSS 4（本地优先，无需后端）"); set_east_asia(r); r.font.size = Pt(11); r.font.color.rgb = MUTED
doc.add_page_break()

# ============ 一、系统概述 ============
add_heading("一、系统概述", 1)
add_para("本系统是一套面向大学物理课程的智慧教学工作台，覆盖学生与教师两类角色。学生端以「过程证据」为主线，把诊断、项目实践、虚拟实验和反思串成一条可复核的学习路径；教师端以「治理」为主线，从班级信号定位风险，并对 AI 建议保留最终复核权。", size=11)
add_para("核心特点：", bold=True)
for t in [
    "闭环学习路径：AI 学情诊断 → 项目匹配 → 五阶任务闯关 → 虚拟实验 + AI 助教 → 学习报告；",
    "过程证据驱动：每一次任务勾选、每一条实验记录、每一轮问答都会写入本地证据并汇入报告；",
    "教师保留最终判断权：AI 只生成建议分，教师复核分与理由始终覆盖建议；",
    "本地优先、可离线：所有数据读写 localStorage，不依赖后端或大模型密钥；",
    "破坏性动作可逆：切换项目、重新诊断、重置记录均需二次确认，并留存 24 小时可恢复快照。",
]:
    p = doc.add_paragraph(t, style="List Bullet")
    for run in p.runs: set_east_asia(run); run.font.size = Pt(11)
doc.add_page_break()

# ============ 二、系统设计图 ============
add_heading("二、系统设计图", 1)

add_heading("2.1 教学闭环原理图", 2)
add_para("下图展示系统如何把学生的每一次学习行动转化为可解释证据：诊断画像驱动项目排序，任务、实验、问答三条证据流汇入同一份学习报告；教师复核覆盖 AI 建议分，教学洞察再把班级信号反哺到下一轮项目匹配。")
add_image(os.path.join(DIAGRAMS, "原理图-教学闭环.png"), 16.5, "图 1 · 教学闭环原理图（学生闭环 + 教师治理 + 可逆保护）")

doc.add_paragraph()
add_heading("2.2 双角色操作流程图", 2)
add_para("下图给出两类角色的完整操作顺序：学生沿「诊断 → 选项目 → 闯关 → 实验与问答 → 报告」五步推进；教师沿「驾驶舱 → 班级画像 → 评价复核 → 教学洞察」完成治理；登录失败与破坏性动作均有明确保护分支。")
add_image(os.path.join(DIAGRAMS, "流程图-双角色操作.png"), 16.5, "图 2 · 双角色操作流程图（含异常与保护分支）")

doc.add_paragraph()
add_heading("2.3 系统分层架构图", 2)
add_para("下图为系统的技术分层：浏览器是唯一运行载体；路由层区分登录页与学生/教师视图；组件层由 physics-app 与 shadcn/ui 构成；逻辑层负责物理计算、报告导出与快捷键；数据层为内置演示数据 + localStorage 持久化 + 24h 撤销快照，全程无需后端。")
add_image(os.path.join(DIAGRAMS, "架构图-系统分层.png"), 16.5, "图 3 · 系统分层架构图（路由 / 组件 / 逻辑 / 数据四层）")
doc.add_page_break()

# ============ 三、功能结构 ============
add_heading("三、功能结构", 1)
table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "角色"; hdr[1].text = "模块"; hdr[2].text = "功能说明"
rows = [
    ("学生", "学习总览", "聚合诊断、项目、任务、评价，给出下一步入口"),
    ("学生", "AI 学情诊断", "8 题前测生成 7 维度知识画像，作为项目匹配起点"),
    ("学生", "章节与知识图谱", "章节 / 概念 / 项目三视角导航课程结构"),
    ("学生", "项目匹配", "按薄弱维度、学习方式与能力目标排序项目案例"),
    ("学生", "任务闯关", "五阶任务：情境认知 → 原理建模 → 方案设计 → 实验验证 → 协同创生"),
    ("学生", "AI 助教知物", "围绕概念、实验与项目提供追问、提示与验证路径"),
    ("学生", "虚拟实验", "超声波测距仿真，调参观察回波时间与误差"),
    ("学生", "学习报告", "汇总诊断、任务、实验与评价，可下载 Markdown / 打印 PDF"),
    ("教师", "教师驾驶舱", "班级趋势、项目进度与即时提醒一览"),
    ("教师", "班级画像", "查看诊断分布、薄弱维度、项目进度与风险状态"),
    ("教师", "章节项目矩阵", "按教材章节查看核心知识、主项目与成果证据"),
    ("教师", "评价复核", "AI 建议分 + 教师复核分，支持单人与批量复核"),
    ("教师", "教学洞察", "基于本周教学信号生成调控方案，并明确 AI 使用边界"),
]
for role, mod, desc in rows:
    c = table.add_row().cells
    c[0].text = role; c[1].text = mod; c[2].text = desc
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            for run in p.runs:
                set_east_asia(run); run.font.size = Pt(10)
doc.add_page_break()

# ============ 四、逐步演示 ============
add_heading("四、逐步操作演示（含截图）", 1)
add_para("以下截图为生产环境（http://localhost:3017）真实操作记录，按「学生线 → 教师线」顺序排列，完整覆盖从登录到闭环的全部关键步骤。", size=10, color=MUTED)

steps = [
    ("4.1 登录页", "01-登录页.png", "输入演示账号（student / teacher）与统一密码 dunzhongwan，点击「进入教学工作台」。登录页使用真实教学场景背景，提亮后校徽与系统名称清晰可见。"),
    ("4.2 学习总览", "02-学习总览.png", "登录后进入学习总览。右侧「学习路径」列出五步完成度；深色 hero 卡给出今天的学习焦点，侧边栏导航对比度已修复，分组标签与菜单项清晰可读。"),
    ("4.3 AI 学情诊断（作答）", "03-AI学情诊断.png", "完成 8 道基础题。右侧实时预览当前维度画像，完成度 8/8 后可提交；提交前题目卡片有明确的选中态与进度指示。"),
    ("4.4 诊断结果", "04-诊断结果.png", "提交后生成诊断结果卡：前测得分 100，完成度 8/8，标记「已保存」。系统用薄弱维度参与后续项目排序，点击「查看项目匹配」进入下一步。"),
    ("4.5 项目匹配", "05-项目匹配.png", "项目匹配中心按匹配度排序，首个卡片标记「AI 优先推荐」（90% 匹配）。点击「选择项目」后卡片出现蓝色选中描边。"),
    ("4.6 任务闯关", "06-任务闯关.png", "进入五阶任务闯关，当前项目为「气体状态参数智能采集与规律拟合」。左侧为五阶步骤条，右侧为当前关卡（第一关：情境认知）。"),
    ("4.7 任务闯关（已保存）", "07-任务闯关-已保存.png", "勾选本关 3 项闯关条件、填写过程说明并保存。总进度变为 20%（1/5 关），第一关步骤条标记为完成，右上角出现「已保存」徽章。"),
    ("4.8 AI 助教知物", "08-AI助教.png", "在 AI 助教页点击快捷问题「为什么超声波可以测距？」，系统以「追问 + 提示」方式回答，不直接给结论，培养先想后答的习惯。"),
    ("4.9 虚拟实验", "09-虚拟实验.png", "超声波测距实验页：左侧为测距场景示意（理论往返时间 5.83 ms），右侧可拖动目标距离、声速、反射率、噪声四个参数。"),
    ("4.10 虚拟实验（测量记录）", "10-虚拟实验-测量记录.png", "点击「运行一次测量」两次，得到测量结果 1.012 m、相对误差 1.17%、信号质量「良好」；下方显示超声脉冲-回波时域信号图。"),
    ("4.11 学习报告", "11-学习报告.png", "学习报告自动汇总：综合得分 75（达标），诊断 100 分、项目进度 20%、实验记录 2 条；支持下载 Markdown / JSON 与打印 PDF。"),
    ("4.12 学习总览（闭环验证）", "12-学习总览-闭环.png", "回到学习总览，学习路径显示 4/5 完成：前测 100 分、已选项目、任务 20%、2 次实验 / 2 条问答，证明数据跨页面贯通。"),
    ("4.13 教师驾驶舱", "13-教师驾驶舱.png", "退出学生账号，以 teacher 登录进入教师驾驶舱。左侧为本周教学调控 hero 卡，右侧即时提醒列出「机械波概念误区」「项目进度滞后」等信号。"),
    ("4.14 班级画像", "14-班级画像.png", "班级画像按风险状态筛选学生，表格展示前测、薄弱维度、当前项目、任务进度与风险等级；李同学（高风险）进度仅 35%。"),
    ("4.15 评价复核", "15-评价复核.png", "评价复核页以卡片形式列出每位学生的四阶能力得分与「最终分 / AI 建议」；支持按风险筛选、全选与批量复核。"),
    ("4.16 评价复核（已保存）", "16-评价复核-已保存.png", "打开张同学复核面板，AI 建议分 73；教师填写复核理由后保存，卡片显示「已复核 · 09/16 12:51」并附上理由摘要。"),
    ("4.17 教学洞察", "17-教学洞察.png", "教学洞察页基于本周信号生成调控方案（如机械波模块的「波速—振动速度」混淆问题），教师可刷新建议、填写补充备注并保存调控记录。"),
]
for title, img, desc in steps:
    add_heading(title, 2)
    add_para(desc)
    add_image(os.path.join(SHOTS, img), 16.5, title)
    doc.add_paragraph()
doc.add_page_break()

# ============ 五、闭环验收结论 ============
add_heading("五、闭环验收结论", 1)
add_para("对照第二章的教学闭环原理图，本次演示已逐项验证以下闭环链路，全部通过：", size=11)
checks = [
    "学生主线闭环：登录 → AI 学情诊断（8/8，100 分）→ 项目匹配（AI 优先推荐）→ 任务闯关（进度 20%）→ AI 助教问答 → 虚拟实验（2 条记录）→ 学习报告（综合 75 分，可下载）；",
    "证据贯通：诊断、任务、实验、问答四类数据均在学习总览与学习报告中一致呈现，跨页面、跨刷新不丢；",
    "教师治理闭环：登录 → 驾驶舱即时提醒 → 班级画像定位高风险学生 → 评价复核（保存复核分与理由）→ 教学洞察生成调控方案；",
    "权限边界：AI 只生成建议分，教师复核结果覆盖建议，最终判断权在教师；",
    "可逆保护：切换项目、重新诊断、重置记录均需二次确认，并自动留存 24 小时可恢复快照；",
    "离线可用：全部数据读写 localStorage，不依赖后端或大模型密钥，课堂断网仍可继续。",
]
for t in checks:
    p = doc.add_paragraph(t, style="List Number")
    for run in p.runs: set_east_asia(run); run.font.size = Pt(11)
doc.add_paragraph()
add_para("结论：系统双角色功能已全部连通，可完成从登录到产出学习报告 / 教学调控记录的完整闭环演示。", bold=True, size=12, color=PRIMARY)

doc.save(OUT)
print("saved:", OUT)
